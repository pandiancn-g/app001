from __future__ import annotations

import math
import random
from datetime import UTC, date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT_PATH = Path("synthetic_data/UC-190_Seasonality_Pattern_Highlighter_Synthetic_Data.xlsx")
SEED = 190


random.seed(SEED)


REGIONS = ["Abu Dhabi City", "Al Ain", "Al Dhafra"]
HOTEL_TYPES = ["Hotel", "Hotel Apartment", "Resort"]
STAR_RATINGS = ["3 Star", "4 Star", "5 Star"]
NATIONALITIES = ["UAE", "India", "UK", "Saudi Arabia", "China", "Germany", "USA", "Egypt"]
VISIT_PURPOSES = ["Leisure", "Business", "Event", "Family Visit"]
MALLS = ["Yas Mall", "Abu Dhabi Mall", "Marina Mall", "Al Wahda Mall", "Galleria Mall"]
PRODUCT_CATEGORIES = ["Fashion", "Electronics", "Dining", "Grocery", "Travel", "Entertainment"]
TRANSPORT_TYPES = ["Bus", "Taxi", "Ferry", "School Bus"]


def daterange(start: date, end: date, step_days: int = 1):
    current = start
    while current <= end:
        yield current
        current += timedelta(days=step_days)


def month_range(start_year: int, start_month: int, end_year: int, end_month: int):
    year = start_year
    month = start_month
    while (year, month) <= (end_year, end_month):
        yield year, month
        month += 1
        if month == 13:
            month = 1
            year += 1


def week_start_dates(start: date, end: date):
    current = start - timedelta(days=start.weekday())
    while current <= end:
        yield current
        current += timedelta(days=7)


def month_end(year: int, month: int) -> date:
    if month == 12:
        return date(year, 12, 31)
    return date(year, month + 1, 1) - timedelta(days=1)


def ramadan_window(year: int) -> tuple[date, date]:
    # Approximate Gregorian dates for synthetic testing only.
    windows = {
        2024: (date(2024, 3, 11), date(2024, 4, 9)),
        2025: (date(2025, 3, 1), date(2025, 3, 29)),
        2026: (date(2026, 2, 18), date(2026, 3, 19)),
        2027: (date(2027, 2, 7), date(2027, 3, 8)),
    }
    return windows.get(year, (date(year, 3, 1), date(year, 3, 29)))


def eid_fitr_window(year: int) -> tuple[date, date]:
    start, end = ramadan_window(year)
    eid_start = end + timedelta(days=1)
    return eid_start, eid_start + timedelta(days=3)


def eid_adha_window(year: int) -> tuple[date, date]:
    # Approximate dates for synthetic testing.
    windows = {
        2024: (date(2024, 6, 16), date(2024, 6, 19)),
        2025: (date(2025, 6, 6), date(2025, 6, 9)),
        2026: (date(2026, 5, 27), date(2026, 5, 30)),
        2027: (date(2027, 5, 16), date(2027, 5, 19)),
    }
    return windows.get(year, (date(year, 6, 10), date(year, 6, 13)))


def is_between(day: date, window: tuple[date, date]) -> bool:
    return window[0] <= day <= window[1]


def count_days_in_window(start: date, end: date, window: tuple[date, date]) -> int:
    return sum(1 for day in daterange(start, end) if is_between(day, window))


def white_friday(year: int) -> date:
    current = date(year, 11, 30)
    while current.weekday() != 4:
        current -= timedelta(days=1)
    return current


def public_holiday_name(day: date) -> str:
    if day.month == 1 and day.day == 1:
        return "New Year"
    if day.month == 12 and day.day in (2, 3):
        return "UAE National Day"
    if is_between(day, eid_fitr_window(day.year)):
        return "Eid Al Fitr"
    if is_between(day, eid_adha_window(day.year)):
        return "Eid Al Adha"
    return ""


def seasonal_features(day: date) -> dict[str, int | str | bool]:
    ramadan = is_between(day, ramadan_window(day.year))
    eid_fitr = is_between(day, eid_fitr_window(day.year))
    eid_adha = is_between(day, eid_adha_window(day.year))
    summer_heat = day.month in (6, 7, 8, 9)
    tourism_peak = day.month in (1, 2, 3, 10, 11, 12)
    retail_event = abs((day - white_friday(day.year)).days) <= 3
    holiday = public_holiday_name(day)
    is_weekend = day.weekday() in (5, 6)
    return {
        "is_weekend": is_weekend,
        "is_working_day": not is_weekend and not holiday,
        "is_ramadan": ramadan,
        "is_eid_al_fitr": eid_fitr,
        "is_eid_al_adha": eid_adha,
        "is_summer_heat": summer_heat,
        "is_tourism_peak": tourism_peak,
        "is_retail_event": retail_event,
        "holiday_name": holiday,
    }


def monthly_feature_counts(year: int, month: int) -> dict[str, int]:
    start = date(year, month, 1)
    end = month_end(year, month)
    counts = {
        "working_days": 0,
        "weekend_days": 0,
        "public_holidays": 0,
        "ramadan_days": 0,
        "eid_days": 0,
        "summer_heat_days": 0,
        "tourism_peak_days": 0,
        "retail_event_days": 0,
    }
    for day in daterange(start, end):
        f = seasonal_features(day)
        counts["working_days"] += int(f["is_working_day"])
        counts["weekend_days"] += int(f["is_weekend"])
        counts["public_holidays"] += int(bool(f["holiday_name"]))
        counts["ramadan_days"] += int(f["is_ramadan"])
        counts["eid_days"] += int(f["is_eid_al_fitr"] or f["is_eid_al_adha"])
        counts["summer_heat_days"] += int(f["is_summer_heat"])
        counts["tourism_peak_days"] += int(f["is_tourism_peak"])
        counts["retail_event_days"] += int(f["is_retail_event"])
    return counts


def safe_round(value: float, digits: int = 2) -> float:
    return round(float(value), digits)


def noise(scale: float) -> float:
    return random.uniform(-scale, scale)


def indicator_components(indicator_id: str, period_date: date, frequency: str, base_index: int = 0):
    month = period_date.month
    year_offset = period_date.year - 2022
    seasonal = 0.0
    calendar = 0.0
    trend = 0.0
    irregular = noise(1.0)
    driver = "None"
    method = "X-13ARIMA-SEATS"

    if frequency == "Monthly":
        counts = monthly_feature_counts(period_date.year, month)
        ramadan_share = counts["ramadan_days"] / max(1, (month_end(period_date.year, month).day))
        eid_share = counts["eid_days"] / max(1, (month_end(period_date.year, month).day))
        heat_share = counts["summer_heat_days"] / max(1, (month_end(period_date.year, month).day))
        tourism_share = counts["tourism_peak_days"] / max(1, (month_end(period_date.year, month).day))
        retail_share = counts["retail_event_days"] / max(1, (month_end(period_date.year, month).day))
        working_days = counts["working_days"]
    else:
        f = seasonal_features(period_date)
        ramadan_share = float(f["is_ramadan"])
        eid_share = float(f["is_eid_al_fitr"] or f["is_eid_al_adha"])
        heat_share = float(f["is_summer_heat"])
        tourism_share = float(f["is_tourism_peak"])
        retail_share = float(f["is_retail_event"])
        working_days = 1 if f["is_working_day"] else 0

    if "HOTEL" in indicator_id or "TOURISM" in indicator_id:
        seasonal = 14 * tourism_share - 9 * heat_share + 5 * eid_share
        calendar = -1.2 if working_days < 20 and frequency == "Monthly" else 0.5 * eid_share
        trend = 1.5 * year_offset + 0.08 * base_index
        driver = "Winter tourism season" if tourism_share else "Summer low season" if heat_share else "Eid travel"
    elif "RETAIL" in indicator_id or "SPEND" in indicator_id or "ECOM" in indicator_id:
        seasonal = 10 * ramadan_share + 13 * eid_share + 9 * retail_share + 4 * tourism_share
        calendar = (working_days - 21) * 0.15 if frequency == "Monthly" else 0.3 * working_days
        trend = 2.0 * year_offset + 0.05 * base_index
        driver = "Ramadan/Eid retail" if ramadan_share or eid_share else "White Friday" if retail_share else "Consumer trend"
    elif "ELECTRICITY" in indicator_id or "WATER" in indicator_id or "TAQA" in indicator_id:
        seasonal = 18 * heat_share + 3 * ramadan_share
        calendar = 0.2 * working_days
        trend = 1.2 * year_offset
        driver = "Summer heat"
        method = "TBATS" if frequency in ("Hourly", "Daily", "Weekly") else "X-13ARIMA-SEATS"
    elif "TRADE" in indicator_id or "CUSTOMS" in indicator_id:
        seasonal = 5 * ramadan_share + 4 * eid_share + 4 * retail_share + 2 * tourism_share
        calendar = (working_days - 21) * 0.3 if frequency == "Monthly" else 0.4 * working_days
        trend = 1.8 * year_offset
        driver = "Trading-day and holiday effect"
    elif "CONSTRUCTION" in indicator_id or "LICENSE" in indicator_id:
        seasonal = -14 * heat_share + 3 * tourism_share
        calendar = (working_days - 21) * 0.2 if frequency == "Monthly" else 0.2 * working_days
        trend = 1.4 * year_offset
        driver = "Summer heat slowdown"
        method = "STL"
    elif "LABOUR" in indicator_id or "UNEMP" in indicator_id or "MOHRE" in indicator_id:
        seasonal = 2.5 * heat_share - 2 * tourism_share + 1.2 * ramadan_share
        calendar = 0.1 * working_days
        trend = -0.15 * year_offset
        driver = "Seasonal hiring cycle"
    elif "TRANSPORT" in indicator_id or "ITC" in indicator_id or "SCHOOL" in indicator_id:
        seasonal = -6 * heat_share + 3 * tourism_share + 2 * ramadan_share
        calendar = 0.25 * working_days
        trend = 0.8 * year_offset
        driver = "School/tourism transport cycle"
        method = "TBATS" if frequency in ("Hourly", "Daily") else "STL"
    else:
        seasonal = 4 * math.sin(2 * math.pi * month / 12)
        trend = 1.0 * year_offset
        driver = "Generic seasonality"

    return seasonal, calendar, trend, irregular, method, driver


def adjusted_record(indicator_id: str, indicator_name: str, period_date: date, frequency: str, base_value: float, row_idx: int):
    seasonal, calendar, trend, irregular, method, driver = indicator_components(indicator_id, period_date, frequency, row_idx)
    raw = base_value + seasonal + calendar + trend + irregular
    adjusted = raw - seasonal - calendar
    seasonal_share = 0 if raw == 0 else abs(seasonal + calendar) / max(abs(raw), 1) * 100
    # Keep this threshold intentionally low enough to create representative
    # warning and alert rows in the synthetic workbook.
    irregular_threshold = abs(irregular) > 0.85
    diagnostics = "Warning" if irregular_threshold else "Pass"
    confidence = "Medium" if diagnostics == "Warning" else "High"
    alert = "Seasonal anomaly" if irregular_threshold else ""
    if method == "STL" and frequency == "Yearly":
        confidence = "Low"
    return {
        "INDICATOR_ID": indicator_id,
        "INDICATOR_NAME": indicator_name,
        "FREQUENCY": frequency,
        "PERIOD_START": period_date.isoformat(),
        "PERIOD_LABEL": period_label(period_date, frequency),
        "RAW_VALUE": safe_round(raw),
        "SEASONAL_COMPONENT": safe_round(seasonal),
        "CALENDAR_COMPONENT": safe_round(calendar),
        "TREND_CYCLE_COMPONENT": safe_round(trend),
        "IRREGULAR_COMPONENT": safe_round(irregular),
        "SEASONALLY_ADJUSTED_VALUE": safe_round(adjusted),
        "SEASONAL_SHARE_PCT": safe_round(seasonal_share),
        "MODEL_METHOD": method,
        "MAIN_SEASONAL_DRIVER": driver,
        "DIAGNOSTIC_STATUS": diagnostics,
        "CONFIDENCE_FLAG": confidence,
        "ALERT_TYPE": alert,
        "PLAIN_LANGUAGE_ANNOTATION": annotation(indicator_name, raw, adjusted, seasonal, calendar, driver, confidence),
    }


def period_label(day: date, frequency: str) -> str:
    if frequency == "Hourly":
        return day.strftime("%Y-%m-%d %H:00")
    if frequency == "Daily":
        return day.isoformat()
    if frequency == "Weekly":
        return f"{day.isocalendar().year}-W{day.isocalendar().week:02d}"
    if frequency == "Monthly":
        return f"{day.year}-{day.month:02d}"
    if frequency == "Yearly":
        return str(day.year)
    return day.isoformat()


def annotation(indicator_name: str, raw: float, adjusted: float, seasonal: float, calendar: float, driver: str, confidence: str) -> str:
    direction = "above" if raw > adjusted else "below"
    return (
        f"{indicator_name} raw value is {direction} the seasonally adjusted estimate. "
        f"The main seasonal driver is {driver}; seasonal plus calendar effects explain "
        f"{safe_round(abs(seasonal + calendar), 1)} points of the movement. "
        f"Confidence is {confidence}; use the adjusted value for short-term momentum analysis."
    )


def add_sheet(wb: Workbook, name: str, rows: list[dict]):
    ws = wb.create_sheet(title=name)
    if not rows:
        return ws
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    style_sheet(ws)
    return ws


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 50))
        ws.column_dimensions[get_column_letter(column)].width = max(12, min(max_length + 2, 42))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def build_readme_rows():
    return [
        {"SECTION": "Workbook", "DETAIL": "UC-190 Seasonality Pattern Highlighter synthetic data"},
        {"SECTION": "Purpose", "DETAIL": "Mock Excel data to test hourly, daily, weekly, monthly, and yearly seasonality detection, decomposition, annotation, publishing, and alert workflows."},
        {"SECTION": "Source alignment", "DETAIL": "Columns are based on attached UC-190 source tables: DCT hotels, AD Customs, TAQA, Mastercard, Visa, DED, ITC, and MOHRE."},
        {"SECTION": "Synthetic note", "DETAIL": "All records are fictional and generated for UI/analytics testing only. Do not treat as official SCAD or partner data."},
        {"SECTION": "Seasonal drivers", "DETAIL": "Ramadan, Eid Al Fitr, Eid Al Adha, summer heat, winter tourism season, working days, White Friday, fiscal cycle, and school transport patterns."},
        {"SECTION": "How to use", "DETAIL": "Use *_OBSERVATIONS sheets for frequency testing, source-specific sheets for schema testing, and SEASONAL_DECOMP_OUTPUT plus REAL_TIME_ALERTS for AI workflow testing."},
    ]


def build_dictionary_rows():
    rows = []
    definitions = [
        ("HOURLY_OBSERVATIONS", "Generic hourly economic indicator feed", "Hourly", "timestamp, indicator_id, raw_value, seasonal flags, model output"),
        ("DAILY_OBSERVATIONS", "Generic daily economic indicator feed", "Daily", "date, indicator_id, raw_value, calendar features"),
        ("WEEKLY_OBSERVATIONS", "Generic weekly economic indicator feed", "Weekly", "week_start, indicator_id, raw_value, seasonal drivers"),
        ("MONTHLY_OBSERVATIONS", "Generic monthly economic indicator feed", "Monthly", "year, month, raw_value, source_table"),
        ("YEARLY_OBSERVATIONS", "Generic yearly economic indicator feed", "Yearly", "year, indicator_id, raw_value"),
        ("DCT_HOTEL_MONTHLY", "DCT hotel/tourism synthetic monthly data", "Monthly", "YEAR, MONTH, OCCUPANCY_RATE, TOTAL_GUESTS, TOTAL_REVENUE, ALOS"),
        ("AD_CUSTOMS_DAILY", "AD Customs foreign trade synthetic daily data", "Daily", "TRANS_ID, BILL_DATE, TRANS_DATE, AMOUNT, WEIGHT_KG, QUANTITY"),
        ("TAQA_MONTHLY", "TAQA electricity/water consumption bills", "Monthly", "CONS_YEAR, CONS_MONTH, QTY, AMT, REGION, TARIFF"),
        ("MC_WEEKLY_SPEND", "Mastercard aggregated spend KPIs", "Weekly", "GROUP_NAME, SEGMENT, CHANNEL, SPEND_CATEGORY, SPEND_AED"),
        ("VISA_HOURLY_RETAIL", "Visa Abu Dhabi retail transactions", "Hourly", "TRANSACTION_DATE, TRANSACTION_TIME, MALL, PRODUCT_CATEGORY, AMOUNT"),
        ("DED_YEARLY_LICENSES", "DED business license synthetic yearly data", "Yearly", "YEAR, LICENSE_TYPE, LICENSE_STATUS, TOTAL_LICENSES"),
        ("ITC_DAILY_TRANSPORT", "ITC transport and school bus synthetic data", "Daily", "TRIP_DATE, TYPE_OF_TRANSPORT, TOTAL_PASSENGERS, TOTAL_STUDENTS"),
        ("MOHRE_MONTHLY_LABOUR", "MOHRE labour synthetic monthly data", "Monthly", "YEAR, MONTH, EMP, CONTRACT_SALARY, AMOUNT_PAID"),
        ("SEASONAL_DECOMP_OUTPUT", "AI/statistical output", "Mixed", "raw_value, seasonal_component, calendar_component, trend, irregular, seasonally_adjusted_value"),
        ("REAL_TIME_ALERTS", "Alerts for analysts", "Mixed", "alert_type, severity, expected_value, actual_value, recommended_action"),
    ]
    for table, purpose, frequency, fields in definitions:
        rows.append({"SHEET_OR_TABLE": table, "PURPOSE": purpose, "FREQUENCY": frequency, "KEY_FIELDS": fields})
    return rows


def build_calendar_rows():
    rows = []
    for day in daterange(date(2025, 1, 1), date(2026, 12, 31)):
        f = seasonal_features(day)
        rows.append({
            "CALENDAR_DATE": day.isoformat(),
            "GREGORIAN_YEAR": day.year,
            "GREGORIAN_MONTH": day.month,
            "DAY_OF_WEEK": day.strftime("%A"),
            "IS_WEEKEND": f["is_weekend"],
            "IS_WORKING_DAY": f["is_working_day"],
            "IS_RAMADAN": f["is_ramadan"],
            "IS_EID_AL_FITR": f["is_eid_al_fitr"],
            "IS_EID_AL_ADHA": f["is_eid_al_adha"],
            "IS_SUMMER_HEAT": f["is_summer_heat"],
            "IS_TOURISM_PEAK": f["is_tourism_peak"],
            "IS_RETAIL_EVENT": f["is_retail_event"],
            "HOLIDAY_NAME": f["holiday_name"],
        })
    return rows


def build_generic_observations(frequency: str):
    configs = [
        ("HOTEL_OCC_RATE", "Hotel Occupancy Rate", "DCT", 70),
        ("RETAIL_SPEND_AED", "Retail Spend AED", "Visa/Mastercard", 120),
        ("ECOM_SPEND_AED", "E-commerce Spend AED", "Visa", 80),
        ("ELECTRICITY_QTY", "Electricity Consumption Quantity", "TAQA", 160),
        ("FOREIGN_TRADE_AMOUNT", "Foreign Trade Amount", "AD Customs", 200),
        ("CONSTRUCTION_OUTPUT", "Construction Output Index", "DED", 95),
        ("LABOUR_EMPLOYMENT", "Employment Count Index", "MOHRE", 105),
        ("TRANSPORT_PASSENGERS", "Transport Passenger Count", "ITC", 85),
    ]
    rows = []
    if frequency == "Hourly":
        periods = [datetime(2026, 2, 16) + timedelta(hours=i) for i in range(24 * 21)]
    elif frequency == "Daily":
        periods = [datetime.combine(day, time()) for day in daterange(date(2026, 1, 1), date(2026, 12, 31))]
    elif frequency == "Weekly":
        periods = [datetime.combine(day, time()) for day in week_start_dates(date(2024, 1, 1), date(2026, 12, 31))]
    elif frequency == "Monthly":
        periods = [datetime(year, month, 1) for year, month in month_range(2022, 1, 2026, 12)]
    else:
        periods = [datetime(year, 1, 1) for year in range(2018, 2027)]

    for idx, period in enumerate(periods):
        for indicator_id, name, source, base in configs:
            period_date = period.date()
            record = adjusted_record(indicator_id, name, period_date, frequency, base, idx)
            record.update({
                "SOURCE_SYSTEM": source,
                "REGION": random.choice(REGIONS),
                "DATA_STATUS": "Final" if period_date < date(2026, 10, 1) else "Preliminary",
                "DATA_VINTAGE": "2026-12-31",
            })
            if frequency == "Hourly":
                record["HOUR"] = period.hour
            rows.append(record)
    return rows


def build_dct_hotel_monthly():
    rows = []
    for year, month in month_range(2023, 1, 2026, 12):
        for region in REGIONS:
            for hotel_type in HOTEL_TYPES:
                base = 62 + (8 if region == "Abu Dhabi City" else 0) + (5 if hotel_type == "Resort" else 0)
                rec = adjusted_record("HOTEL_OCC_RATE", "Hotel Occupancy Rate", date(year, month, 1), "Monthly", base, month)
                occupancy = min(96, max(35, rec["RAW_VALUE"]))
                total_rooms = random.randint(1800, 5200)
                occupied_rooms = int(total_rooms * occupancy / 100)
                total_guests = int(occupied_rooms * random.uniform(1.5, 2.4))
                alos = random.uniform(2.0, 4.8) + (0.6 if month in (1, 2, 3, 11, 12) else 0)
                guest_nights = int(total_guests * alos)
                room_revenue = occupied_rooms * random.uniform(330, 820)
                fb_revenue = room_revenue * random.uniform(0.18, 0.35)
                other_revenue = room_revenue * random.uniform(0.05, 0.16)
                rows.append({
                    "TECH_TABLE_NAME": "R_E01_D01_HOTEL_STATS",
                    "YEAR": year,
                    "MONTH": month,
                    "CITY": region,
                    "HOTEL_TYPE": hotel_type,
                    "HOTEL_RATING": random.choice(STAR_RATINGS),
                    "NUMBER_OF_HOTELS": random.randint(10, 80),
                    "NUMBER_OF_ROOMS": total_rooms,
                    "TOTAL_GUESTS": total_guests,
                    "TOTAL_GUEST_NIGHTS": guest_nights,
                    "AVG_LENGH_OF_STAY": safe_round(alos),
                    "OCCUPANCY_RATE": safe_round(occupancy),
                    "ROOM_REVENUE": safe_round(room_revenue),
                    "FOOD_BEVERAGE_REVENUE": safe_round(fb_revenue),
                    "OTHER_REVENUE": safe_round(other_revenue),
                    "TOTAL_REVENUE": safe_round(room_revenue + fb_revenue + other_revenue),
                    "SEASONAL_DRIVER": rec["MAIN_SEASONAL_DRIVER"],
                })
    return rows


def build_ad_customs_daily():
    rows = []
    trans_id = 100000
    for day in daterange(date(2026, 1, 1), date(2026, 6, 30)):
        f = seasonal_features(day)
        daily_count = 6 + int(f["is_working_day"]) * 8 + int(f["is_retail_event"]) * 5 + int(f["is_ramadan"]) * 2
        for _ in range(daily_count):
            amount = random.uniform(80000, 1800000) * (1.18 if f["is_ramadan"] or f["is_retail_event"] else 1)
            quantity = random.randint(20, 4000)
            rows.append({
                "TECH_TABLE_NAME": "R_E04_D01_GETFORIEGNTRADE",
                "TRANS_ID": trans_id,
                "CENTER_CODE": random.choice(["AUH_PORT", "KIZAD", "AIRPORT", "AL_AIN"]),
                "CENTER_NAME": random.choice(["Abu Dhabi Port", "Khalifa Port", "Abu Dhabi Airport", "Al Ain Customs"]),
                "CENTER_TYPE": random.choice(["Sea", "Air", "Land"]),
                "BILL_TYPE": random.choice(["Import", "Export", "Re-export"]),
                "BILL_DATE": day.isoformat(),
                "HARM_CODE": random.choice(["851712", "870323", "100630", "271019", "620342"]),
                "HARM_DESC_E": random.choice(["Mobile phones", "Motor vehicles", "Rice", "Petroleum oils", "Garments"]),
                "ORIGIN_COUNTRY": random.choice(["China", "India", "USA", "Germany", "Saudi Arabia"]),
                "DEST_COUNTRY": random.choice(["UAE", "Saudi Arabia", "India", "Oman", "Kuwait"]),
                "UNIT_OF_MEASURE": random.choice(["KG", "PCS", "LTR"]),
                "AMOUNT": safe_round(amount),
                "WEIGHT_KG": safe_round(quantity * random.uniform(0.8, 18.0)),
                "QUANTITY": quantity,
                "DUTY_AMOUNT": safe_round(amount * random.uniform(0.01, 0.05)),
                "EMARA": "Abu Dhabi",
                "TRANS_DATE": day.isoformat(),
                "SEASONAL_DRIVER": "Ramadan import surge" if f["is_ramadan"] else "Retail event import surge" if f["is_retail_event"] else "Working-day trade flow",
            })
            trans_id += 1
    return rows


def build_taqa_monthly():
    rows = []
    for year, month in month_range(2023, 1, 2026, 12):
        for region in REGIONS:
            for premise_type in ["Villa", "Apartment", "Commercial", "Hotel"]:
                base = 2600 if premise_type in ("Villa", "Hotel") else 1500
                rec = adjusted_record("TAQA_ELECTRICITY", "Electricity Consumption", date(year, month, 1), "Monthly", base, month)
                qty = max(300, rec["RAW_VALUE"] * random.uniform(1.0, 1.4))
                amt = qty * random.uniform(0.23, 0.38)
                rows.append({
                    "TECH_TABLE_NAME": "R_E14_D01_DISTRIBUTION_COMP",
                    "CONS_YEAR": year,
                    "CONS_MONTH": month,
                    "CONS_DAYS": month_end(year, month).day,
                    "REGION": region,
                    "REGION_DESC": region,
                    "PREMISE_TYPE": premise_type,
                    "TARIFF": random.choice(["Residential UAE National", "Residential Expat", "Commercial", "Hotel"]),
                    "QTY": safe_round(qty),
                    "AMT": safe_round(amt),
                    "AVG_DAILY_E_CONS_QTY": safe_round(qty / month_end(year, month).day),
                    "AVG_DAILY_E_CONS_AMT": safe_round(amt / month_end(year, month).day),
                    "SA_STATUS": "Active",
                    "CITY": region,
                    "SEASONAL_DRIVER": rec["MAIN_SEASONAL_DRIVER"],
                })
    return rows


def build_mastercard_weekly():
    rows = []
    for week in week_start_dates(date(2025, 1, 1), date(2026, 12, 31)):
        for category in ["Dining", "Travel", "Retail", "Grocery", "Hotels"]:
            rec = adjusted_record("MC_SPEND_AED", "Mastercard Weekly Spend", week, "Weekly", 500000, week.isocalendar().week)
            multiplier = 1.4 if category in ("Retail", "Grocery") and seasonal_features(week)["is_ramadan"] else 1.0
            spend = max(10000, rec["RAW_VALUE"] * 1000 * multiplier)
            transactions = int(spend / random.uniform(120, 450))
            rows.append({
                "TECH_TABLE_NAME": "R_E105_D01_MASTERCARD_AGGREGATE_SPENDING",
                "WEEK_START_DATE": week.isoformat(),
                "YEAR": week.isocalendar().year,
                "WEEK": week.isocalendar().week,
                "GROUP_NAME": random.choice(["Group_A", "Group_B", "Group_C", "Group_D"]),
                "CR_DR_PREPAID": random.choice(["Credit", "Debit", "Prepaid"]),
                "SEGMENT": random.choice(["Consumer", "Commercial"]),
                "CUSTOMER_CLASS": random.choice(["Mass", "Affluent", "High Affluent"]),
                "CARD_SOURCE": random.choice(["Domestic", "International"]),
                "SPEND_COUNTRY": "UAE",
                "SPEND_COUNTRY_AD_SPLIT": "Abu Dhabi",
                "CHANNEL": random.choice(["Online", "Offline"]),
                "MCC_CLASS_NAME": category,
                "SPEND_CATEGORY": category,
                "SPEND_CATEGORIZATION": random.choice(["Essential", "Discretionary", "Travel"]),
                "SPEND_TYPE": random.choice(["POS", "E-commerce"]),
                "SPEND_AED": safe_round(spend),
                "SPC": safe_round(spend / max(1, transactions)),
                "TRANSACTIONS": transactions,
                "SEASONAL_DRIVER": rec["MAIN_SEASONAL_DRIVER"],
            })
    return rows


def build_visa_hourly():
    rows = []
    tx_id = 900000
    start_dt = datetime(2026, 2, 16)
    for hour_idx in range(24 * 21):
        ts = start_dt + timedelta(hours=hour_idx)
        f = seasonal_features(ts.date())
        hour_factor = 1.0
        if f["is_ramadan"] and ts.hour in (20, 21, 22, 23):
            hour_factor = 2.4
        elif ts.hour in (12, 13, 18, 19):
            hour_factor = 1.35
        for mall in MALLS:
            for category in random.sample(PRODUCT_CATEGORIES, 3):
                base_amount = random.uniform(800, 12000) * hour_factor
                rows.append({
                    "TECH_TABLE_NAME": "R_E45_D04_VISA_ABUDHABI_RETAIL",
                    "TRANSACTION_DATE": ts.date().isoformat(),
                    "TRANSACTION_TIME": ts.strftime("%H:%M:%S"),
                    "TRANSACTION_ID": tx_id,
                    "MALL": mall,
                    "PRODUCT_CATEGORY": category,
                    "STORE_CATEGORY": random.choice(["Anchor", "Specialty", "Kiosk", "Food Court"]),
                    "TYPE_OF_PRODUCT": random.choice(["Goods", "Service", "Food and Beverage"]),
                    "EMIRATE_RESIDENTSHIP": random.choice(["Abu Dhabi Resident", "Other UAE", "Visitor"]),
                    "SHOP_NAME": f"{category} Store {random.randint(1, 99)}",
                    "NEW_SHOP_STATUS": random.choice(["Existing", "New"]),
                    "TOTAL_SPACE": random.randint(50, 1400),
                    "START_DATE_OF_SHOP": date(random.randint(2015, 2025), random.randint(1, 12), random.randint(1, 28)).isoformat(),
                    "AMOUNT": safe_round(base_amount),
                    "SEASONAL_DRIVER": "Ramadan late-night spend" if f["is_ramadan"] and ts.hour >= 20 else "Hourly retail pattern",
                })
                tx_id += 1
    return rows


def build_ded_yearly():
    rows = []
    for year in range(2018, 2027):
        for license_type in ["Commercial", "Industrial", "Tourism", "Professional", "E-commerce"]:
            base = 1200 + (year - 2018) * 95
            if license_type == "E-commerce":
                base += (year - 2018) * 80
            rows.append({
                "TECH_TABLE_NAME": "R_E03_D04_LICENSE",
                "POST_YEAR": year,
                "YEAR": year,
                "LICENSE_TYPE_EN": license_type,
                "LICENSE_STATUS_EN": "Active",
                "LICENSE_ISSUE_PLACE_ENU": random.choice(REGIONS),
                "FREE_ZONE_NAME_ENU": random.choice(["None", "ADGM", "KIZAD", "twofour54"]),
                "TOTAL_LICENSES": int(base + noise(80)),
                "NEW_LICENSES": int(base * random.uniform(0.12, 0.22)),
                "RENEWED_LICENSES": int(base * random.uniform(0.55, 0.75)),
                "CANCELLED_LICENSES": int(base * random.uniform(0.03, 0.08)),
                "FISCAL_YEAR_CYCLE_FLAG": year >= 2024,
                "SEASONAL_DRIVER": "Fiscal year cycle / annual business registration",
            })
    return rows


def build_itc_daily():
    rows = []
    for day in daterange(date(2026, 1, 1), date(2026, 12, 31)):
        f = seasonal_features(day)
        for transport_type in TRANSPORT_TYPES:
            base = 45000 if transport_type == "Bus" else 18000 if transport_type == "Taxi" else 2500
            if transport_type == "School Bus":
                base = 12000
                if day.month in (7, 8) or f["is_weekend"]:
                    base *= 0.15
            value = base * (1.15 if f["is_tourism_peak"] else 1) * (0.8 if f["is_summer_heat"] else 1) + noise(base * 0.08)
            rows.append({
                "TECH_TABLE_NAME": "R_E86_D08_ITC_HAFILAT_PASSENGERS",
                "TRIP_DATE": day.isoformat(),
                "TYPE_OF_TRANSPORT": transport_type,
                "REGION": random.choice(REGIONS),
                "TOTAL_PASSENGERS": max(0, int(value)),
                "TOTAL_STUDENTS": max(0, int(value * 0.65)) if transport_type == "School Bus" else 0,
                "STARTED_TIME": "07:00:00" if transport_type == "School Bus" else "00:00:00",
                "COMPLETED_TIME": "16:00:00" if transport_type == "School Bus" else "23:59:00",
                "DIRECTION": random.choice(["Inbound", "Outbound", "Both"]),
                "SEASONAL_DRIVER": "School calendar" if transport_type == "School Bus" else "Tourism and working-day transport cycle",
            })
    return rows


def build_mohre_monthly():
    rows = []
    for year, month in month_range(2023, 1, 2026, 12):
        for sector in ["Construction", "Hospitality", "Retail", "Transport", "Manufacturing"]:
            base = 80000 if sector == "Construction" else 45000
            rec = adjusted_record("MOHRE_LABOUR_EMP", "MOHRE Employment", date(year, month, 1), "Monthly", base, month)
            emp = max(5000, int(rec["RAW_VALUE"]))
            salary = random.uniform(2800, 12500)
            rows.append({
                "TECH_TABLE_NAME": "R_E08_D01_EMPLOYEE",
                "DATE_TIME_KEY": f"{year}{month:02d}",
                "YEAR": year,
                "MONTH": month,
                "COMPANY_SECTION_ID": sector,
                "EMP": emp,
                "CONTRACT_SALARY": safe_round(salary),
                "AMOUNT_PAID": safe_round(salary * random.uniform(0.96, 1.03)),
                "EMP_GENDER_ID": random.choice(["Male", "Female"]),
                "EMP_NATIONALITY_UNSD": random.choice(["356", "586", "818", "784", "608"]),
                "EMP_AGE_CLASS_ID": random.choice(["18-29", "30-44", "45-59", "60+"]),
                "JOB_ENSCO_L6_ID": random.choice(["Service", "Professional", "Technician", "Craft", "Manager"]),
                "SEASONAL_DRIVER": rec["MAIN_SEASONAL_DRIVER"],
            })
    return rows


def build_model_specs():
    specs = [
        ("HOTEL_OCC_RATE", "Hotel Occupancy Rate", "X-13ARIMA-SEATS", "tourism_peak,event_calendar,eid", "Monthly"),
        ("RETAIL_SPEND_AED", "Retail Spend AED", "X-13ARIMA-SEATS", "ramadan,eid,white_friday,working_days", "Monthly"),
        ("VISA_HOURLY_RETAIL", "Visa Hourly Retail", "TBATS", "hour_of_day,ramadan,eid,weekend,white_friday", "Hourly"),
        ("ELECTRICITY_QTY", "Electricity Consumption", "TBATS", "temperature,summer_heat,weekday", "Daily"),
        ("FOREIGN_TRADE_AMOUNT", "Foreign Trade Amount", "X-13ARIMA-SEATS", "working_days,ramadan,eid,white_friday", "Daily"),
        ("CONSTRUCTION_OUTPUT", "Construction Output Index", "STL", "summer_heat,working_days", "Monthly"),
        ("LABOUR_EMPLOYMENT", "Employment Count Index", "X-13ARIMA-SEATS", "tourism_hiring,summer_slowdown,ramadan", "Monthly"),
        ("TRANSPORT_PASSENGERS", "Transport Passenger Count", "TBATS", "weekday,school_calendar,tourism_peak,ramadan", "Daily"),
    ]
    rows = []
    for idx, (indicator_id, name, method, regressors, freq) in enumerate(specs, start=1):
        rows.append({
            "MODEL_SPEC_ID": f"UC190_MS_{idx:03d}",
            "INDICATOR_ID": indicator_id,
            "INDICATOR_NAME": name,
            "FREQUENCY": freq,
            "MODEL_TYPE": method,
            "DECOMPOSITION_TYPE": "Auto additive/multiplicative",
            "CALENDAR_VARIABLES": regressors,
            "OUTLIER_DETECTION": True,
            "MINIMUM_HISTORY_PERIODS": 84 if freq == "Monthly" else 365,
            "VERSION": "v1.0-synthetic",
            "APPROVAL_STATUS": "Research mock",
        })
    return rows


def build_alerts(decomp_rows):
    candidates = [row for row in decomp_rows if row["DIAGNOSTIC_STATUS"] == "Warning"][:20]
    rows = []
    for idx, row in enumerate(candidates, start=1):
        severity = "High" if row["CONFIDENCE_FLAG"] == "Low" else random.choice(["Medium", "Advisory"])
        rows.append({
            "ALERT_ID": f"UC190_ALERT_{idx:04d}",
            "INDICATOR_ID": row["INDICATOR_ID"],
            "INDICATOR_NAME": row["INDICATOR_NAME"],
            "FREQUENCY": row["FREQUENCY"],
            "PERIOD_LABEL": row["PERIOD_LABEL"],
            "ALERT_TYPE": random.choice(["Seasonal anomaly", "Residual seasonality", "Evolving seasonality", "Short history warning"]),
            "SEVERITY": severity,
            "EXPECTED_VALUE": safe_round(row["SEASONALLY_ADJUSTED_VALUE"]),
            "ACTUAL_VALUE": safe_round(row["RAW_VALUE"]),
            "DEVIATION": safe_round(row["RAW_VALUE"] - row["SEASONALLY_ADJUSTED_VALUE"]),
            "MAIN_SEASONAL_DRIVER": row["MAIN_SEASONAL_DRIVER"],
            "RECOMMENDED_ACTION": "Review model diagnostics and validate with domain specialist before publication.",
            "STATUS": random.choice(["Open", "Under Review", "Closed"]),
            "CREATED_AT": datetime(2026, 12, 31, 9, idx % 60).isoformat(),
        })
    return rows


def build_workbook():
    wb = Workbook()
    wb.remove(wb.active)

    add_sheet(wb, "00_README", build_readme_rows())
    add_sheet(wb, "01_DATA_DICTIONARY", build_dictionary_rows())
    add_sheet(wb, "02_CALENDAR_FEATURES", build_calendar_rows())

    hourly = build_generic_observations("Hourly")
    daily = build_generic_observations("Daily")
    weekly = build_generic_observations("Weekly")
    monthly = build_generic_observations("Monthly")
    yearly = build_generic_observations("Yearly")

    add_sheet(wb, "HOURLY_OBSERVATIONS", hourly)
    add_sheet(wb, "DAILY_OBSERVATIONS", daily)
    add_sheet(wb, "WEEKLY_OBSERVATIONS", weekly)
    add_sheet(wb, "MONTHLY_OBSERVATIONS", monthly)
    add_sheet(wb, "YEARLY_OBSERVATIONS", yearly)

    add_sheet(wb, "DCT_HOTEL_MONTHLY", build_dct_hotel_monthly())
    add_sheet(wb, "AD_CUSTOMS_DAILY", build_ad_customs_daily())
    add_sheet(wb, "TAQA_MONTHLY", build_taqa_monthly())
    add_sheet(wb, "MC_WEEKLY_SPEND", build_mastercard_weekly())
    add_sheet(wb, "VISA_HOURLY_RETAIL", build_visa_hourly())
    add_sheet(wb, "DED_YEARLY_LICENSES", build_ded_yearly())
    add_sheet(wb, "ITC_DAILY_TRANSPORT", build_itc_daily())
    add_sheet(wb, "MOHRE_MONTHLY_LABOUR", build_mohre_monthly())

    decomp = []
    for rows in (hourly[:600], daily[:900], weekly[:600], monthly, yearly):
        decomp.extend(rows)
    add_sheet(wb, "SEASONAL_DECOMP_OUTPUT", decomp)
    add_sheet(wb, "MODEL_SPECS", build_model_specs())
    add_sheet(wb, "REAL_TIME_ALERTS", build_alerts(decomp))

    dashboard_rows = [
        {"METRIC": "Frequencies covered", "VALUE": "Hourly, Daily, Weekly, Monthly, Yearly"},
        {"METRIC": "Source domains covered", "VALUE": "DCT, AD Customs, TAQA, Mastercard, Visa, DED, ITC, MOHRE"},
        {"METRIC": "Seasonal drivers", "VALUE": "Ramadan, Eid, summer heat, tourism peak, working days, White Friday, fiscal cycle"},
        {"METRIC": "Model methods", "VALUE": "X-13ARIMA-SEATS, STL, TBATS"},
        {"METRIC": "Synthetic workbook generated at", "VALUE": datetime.now(UTC).replace(microsecond=0).isoformat()},
    ]
    add_sheet(wb, "DASHBOARD_SAMPLE", dashboard_rows)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_workbook()
    print(path)
